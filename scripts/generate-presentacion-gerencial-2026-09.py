#!/usr/bin/env python3
"""Genera presentación gerencial septiembre 2026 (máx. 4 diapositivas)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "docs" / "informes" / "presentacion-gerencial-statflow-2026-09.pptx"
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
GRAY = RGBColor(0x40, 0x40, 0x40)


def set_title(slide, title: str, subtitle: str | None = None):
    slide.shapes.title.text = title
    tf = slide.shapes.title.text_frame
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = NAVY
    if subtitle and len(slide.placeholders) > 1:
        ph = slide.placeholders[1]
        ph.text = subtitle
        for p in ph.text_frame.paragraphs:
            p.font.size = Pt(16)
            p.font.color.rgb = GRAY


def add_bullets(slide, items: list[str], left=0.6, top=1.6, width=8.8, height=5.0, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = GRAY
        p.space_after = Pt(10)


prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Slide 1 — Portada + mensaje clave
s1 = prs.slides.add_slide(prs.slide_layouts[0])
set_title(
    s1,
    "SJ StatFlow — Avance septiembre 2026",
    "Informe gerencial · SJ Seguridad",
)
if len(s1.placeholders) > 1:
    s1.placeholders[1].text_frame.add_paragraph().text = (
        "GH digitalizada: desvinculaciones, cursos, selección y acreditaciones con validación APO"
    )
    for p in s1.placeholders[1].text_frame.paragraphs:
        p.font.size = Pt(14)
        p.font.color.rgb = BLUE

# Slide 2 — GH
s2 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
set_title(s2, "Gestión Humana — entregas del mes")
add_bullets(
    s2,
    [
        "Desvinculaciones: masivos, ZIP de cartas, seguimientos con checklist OK TODO",
        "Cursos: vigencia automática, catálogo, import/export Excel, documentos por registro",
        "Selección: dashboard KPIs, ingreso y examen ocupacional, export y filtros",
        "Acreditaciones: acreditados, reporte diario (2 Excel APO), 4 colas de validación cruzada",
        "En curso: Export SuperVigilancia (.xls) + dashboard acreditaciones",
    ],
    top=1.4,
    size=17,
)

# Slide 3 — Transversal + indicadores
s3 = prs.slides.add_slide(prs.slide_layouts[5])
set_title(s3, "Transversal, TIC e indicadores")
left_col = [
    "TIC: solicitudes de desarrollo FO-TIC-23 (radicación, líder, bandeja, chat)",
    "Requisiciones: impresión y notificaciones ampliadas",
    "Compras: comentarios y plantilla masiva de precarga",
    "Ficha empleados: campos, import y permisos copy/edit",
    "UX: iconos unificados, marca corporativa en tableros",
]
right_box = s3.shapes.add_textbox(Inches(0.6), Inches(1.35), Inches(5.2), Inches(4.8))
tf = right_box.text_frame
for i, t in enumerate(left_col):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = "• " + t
    p.font.size = Pt(15)
    p.font.color.rgb = GRAY
    p.space_after = Pt(8)

kpi = s3.shapes.add_shape(
    1, Inches(6.0), Inches(1.35), Inches(3.4), Inches(4.8)
)  # rectangle
kpi.fill.solid()
kpi.fill.fore_color.rgb = RGBColor(0xE8, 0xF0, 0xF8)
kpi.line.color.rgb = BLUE
ktf = kpi.text_frame
ktf.word_wrap = True
ktf.paragraphs[0].text = "Indicadores sep."
ktf.paragraphs[0].font.bold = True
ktf.paragraphs[0].font.size = Pt(18)
ktf.paragraphs[0].font.color.rgb = NAVY
for line in [
    "37 commits",
    "5 features cerradas",
    "4 en cierre / activas",
    "Cierre: Acreditaciones operativas",
]:
    p = ktf.add_paragraph()
    p.text = line
    p.font.size = Pt(16)
    p.font.color.rgb = GRAY
    p.space_before = Pt(12)

# Slide 4 — Próximos pasos
s4 = prs.slides.add_slide(prs.slide_layouts[5])
set_title(s4, "Próximos pasos y recomendaciones")
add_bullets(
    s4,
    [
        "Cerrar Cursos, TIC y cola «sin curso» (documentación y permisos)",
        "Completar Export APO + dashboard; prueba con lote real SuperVigilancia",
        "Capacitar GH: rutina diaria Reporte Diario → Validaciones → acciones",
        "Definir matriz rol–permiso antes de adopción masiva de tableros",
    ],
    top=1.4,
    size=19,
)

prs.save(OUT)
print(f"Generado: {OUT}")
